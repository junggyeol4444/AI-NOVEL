"""Stage 15 DOCX parser tests; all sample text is synthetic and ephemeral."""
from __future__ import annotations

from io import BytesIO
from pathlib import Path
import hashlib
import tempfile
import unittest
import warnings
from unittest.mock import patch
from zipfile import ZipFile, ZipInfo

from ai_novel_factory.references import ReferenceParseError, parse_docx_bytes, parse_docx_file

try:
    from docx import Document
except ImportError:
    Document = None


def _fixture(*, table: bool = False, header: bool = False) -> bytes:
    doc = Document()
    doc.add_paragraph("첫 번째 단락")
    doc.add_heading("가상 제목", level=2)
    if table:
        grid = doc.add_table(rows=2, cols=2)
        grid.cell(0, 0).text = "A1"
        grid.cell(0, 1).text = "B1"
        grid.cell(1, 0).text = "A2"
        grid.cell(1, 1).text = "B2"
    doc.add_paragraph("마지막 단락")
    if header:
        doc.sections[0].header.paragraphs[0].text = "비공개 머리글"
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def _repack(raw: bytes, *, extra_name: str, extra_content: bytes) -> bytes:
    dest = BytesIO()
    with ZipFile(BytesIO(raw)) as source, ZipFile(dest, "w") as target:
        for item in source.infolist():
            target.writestr(item, source.read(item.filename))
        target.writestr(extra_name, extra_content)
    return dest.getvalue()


@unittest.skipUnless(Document is not None, "python-docx optional dependency missing")
class DocxParserTests(unittest.TestCase):
    def test_paragraphs_in_original_order(self):
        result = parse_docx_bytes(_fixture())
        self.assertEqual(result.format, "docx")
        self.assertEqual([b.text for b in result.blocks], ["첫 번째 단락", "가상 제목", "마지막 단락"])
        self.assertEqual([b.kind for b in result.blocks], ["paragraph", "heading", "paragraph"])
        self.assertEqual(result.blocks[1].heading_level, 2)
        self.assertEqual(result.paragraph_count, 2)

    def test_table_between_paragraphs(self):
        result = parse_docx_bytes(_fixture(table=True))
        self.assertEqual([b.kind for b in result.blocks], ["paragraph", "heading", "table", "paragraph"])
        self.assertEqual(result.blocks[2].text, "A1\tB1\nA2\tB2")
        self.assertEqual([b.start_line for b in result.blocks], [1, 2, 3, 4])

    def test_header_not_extracted(self):
        result = parse_docx_bytes(_fixture(header=True))
        self.assertNotIn("비공개 머리글", result.text)

    def test_no_original_text_in_repr(self):
        result = parse_docx_bytes(_fixture())
        self.assertNotIn("첫 번째 단락", repr(result))
        self.assertNotIn("첫 번째 단락", repr(result.blocks[0]))

    def test_file_is_read_only_and_unchanged(self):
        raw = _fixture()
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "sample.docx"
            path.write_bytes(raw)
            before = hashlib.sha256(path.read_bytes()).hexdigest()
            result = parse_docx_file(path)
            self.assertEqual(result.blocks[0].text, "첫 번째 단락")
            self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), before)

    def test_invalid_extension_rejected(self):
        with self.assertRaisesRegex(ReferenceParseError, "DOCX 확장자"):
            parse_docx_file("something.docm")

    def test_symlink_rejected(self):
        with tempfile.TemporaryDirectory() as folder:
            target = Path(folder) / "original.docx"
            target.write_bytes(_fixture())
            link = Path(folder) / "link.docx"
            try:
                link.symlink_to(target)
            except (OSError, NotImplementedError):
                self.skipTest("symbolic links not permitted on this system")
            with self.assertRaisesRegex(ReferenceParseError, "심볼릭"):
                parse_docx_file(link)

    def test_corrupt_zip_rejected(self):
        with self.assertRaises(ReferenceParseError):
            parse_docx_bytes(b"not a docx")

    def test_missing_required_part_rejected(self):
        stream = BytesIO()
        with ZipFile(stream, "w") as archive:
            archive.writestr("[Content_Types].xml", b"<Types/>")
        with self.assertRaisesRegex(ReferenceParseError, "필수 구성요소"):
            parse_docx_bytes(stream.getvalue())

    def test_traversal_entry_rejected(self):
        raw = _repack(_fixture(), extra_name="../malicious.txt", extra_content=b"test")
        with self.assertRaisesRegex(ReferenceParseError, "경로"):
            parse_docx_bytes(raw)

    def test_duplicate_member_rejected(self):
        raw = _fixture()
        with ZipFile(BytesIO(raw)) as z:
            name = z.namelist()[0]
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)
            changed = _repack(raw, extra_name=name, extra_content=b"again")
        with self.assertRaisesRegex(ReferenceParseError, "경로"):
            parse_docx_bytes(changed)

    def test_macro_entry_rejected(self):
        raw = _repack(_fixture(), extra_name="word/vbaProject.bin", extra_content=b"dummy")
        with self.assertRaisesRegex(ReferenceParseError, "매크로"):
            parse_docx_bytes(raw)

    def test_oversize_text_rejected(self):
        with self.assertRaisesRegex(ReferenceParseError, "본문 길이"):
            parse_docx_bytes(_fixture(), max_text_chars=3)

    def test_empty_body_rejected(self):
        document = Document()
        stream = BytesIO()
        document.save(stream)
        with self.assertRaisesRegex(ReferenceParseError, "비어"):
            parse_docx_bytes(stream.getvalue())

    def test_oversize_file_rejected(self):
        raw = _fixture()
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "oversize.docx"
            path.write_bytes(raw)
            with self.assertRaisesRegex(ReferenceParseError, "허용 크기"):
                parse_docx_file(path, max_bytes=8)

    def test_uncompressed_member_limit_rejected(self):
        with patch("ai_novel_factory.references.docx_parser.MAX_MEMBER_BYTES", 4):
            with self.assertRaisesRegex(ReferenceParseError, "내부 항목"):
                parse_docx_bytes(_fixture())

    def test_total_uncompressed_limit_rejected(self):
        with patch("ai_novel_factory.references.docx_parser.MAX_UNCOMPRESSED_BYTES", 4):
            with self.assertRaisesRegex(ReferenceParseError, "압축 해제 크기"):
                parse_docx_bytes(_fixture())

    def test_unsafe_limits_rejected(self):
        with self.assertRaisesRegex(ReferenceParseError, "본문 길이 제한"):
            parse_docx_bytes(_fixture(), max_text_chars=True)
        with self.assertRaisesRegex(ReferenceParseError, "파일 크기 제한"):
            parse_docx_file("fake.docx", max_bytes=-1)


if __name__ == "__main__":
    unittest.main()