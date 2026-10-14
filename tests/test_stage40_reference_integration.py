"""Stage 40: offline integration from synthetic reference files to safety gates.

All source paragraphs are invented in this test. Test fixtures are created in
TemporaryDirectory, never shipped as reference originals or persisted in SQLite.
"""
from __future__ import annotations

from contextlib import redirect_stderr, redirect_stdout
from io import BytesIO, StringIO
import hashlib
from pathlib import Path
import shutil
import socket
import tempfile
import unittest
from unittest.mock import patch

from ai_novel_factory.database import ReferenceCatalog
from ai_novel_factory.database.element_selections import ElementSelectionStore, ElementSelectionStorageError
from ai_novel_factory.database.genre_patterns import GenrePatternStore, GenreStorageError
from ai_novel_factory.database.pattern_strengths import PatternStrengthStore
from ai_novel_factory.database.pattern_usage import PatternUsageError, PatternUsageStore
from ai_novel_factory.database.reference_profiles import ReferenceProfileStore
from ai_novel_factory.references.plaintext_parser import ReferenceParseError
from ai_novel_factory.references.profile_cli import main as profile_cli
from ai_novel_factory.references.profile_pipeline import import_reference_profile
from ai_novel_factory.references.related_pattern_search import RelatedPatternSearch, RelatedPatternError
from ai_novel_factory.references.reference_profile import render_reference_profile, validate_profile
from ai_novel_factory.references.similarity_baseline import (
    analyze_similarity_baseline, require_similarity_safe_draft, SimilarityCheckError,
)
from ai_novel_factory.references.similarity_cli import main as similarity_cli

try:
    from docx import Document
except ImportError:
    Document = None
try:
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
except ImportError:
    PdfWriter = None

CANARY = "SYNTHETIC_PRIVATE_BODY_NEVER_PERSIST_20261012"
NEAR_SOURCE = "어두운 골목 끝에서 길을 잃은 사람이 오래된 지도를 발견하고 집으로 돌아갈 길을 찾아냈다."
UNRELATED = "아침 햇살이 창가를 지나가자 작은 새들이 노래했고 사람들은 공원에서 운동을 시작했다."
FANTASY = ("그는 어둠 속에서 지도를 펼쳤다. 그리고 마침내 비밀을 발견했다. ",
           "그는 산을 넘어 새로운 성을 보았다. 그러나 갑자기 문이 열렸다. ")
ROMANCE = ("그녀는 꽃을 보고 기뻤다. 따뜻한 햇살이 반짝였다. ",
           "친구들은 함께 웃었다. 사랑과 희망을 나누었다. ")


def _synthetic_text(genre: str, index: int) -> str:
    phrases = FANTASY if genre == "fantasy" else ROMANCE
    marker = "# " if genre == "romance" else ""
    return (f"{marker}제1화\n\n{phrases[0] * (7 + index)}\n\n"
            f"{marker}제2화\n\n{phrases[1] * (7 + index)}\n\n{CANARY}")


class ReferenceAnalysisIntegrationTests(unittest.TestCase):
    """Real on-disk files, real SQLite, real import/analysis/pattern components."""

    @classmethod
    def setUpClass(cls):
        cls.root = tempfile.TemporaryDirectory()
        cls.base = Path(cls.root.name)
        cls.db_template = cls.base / "template.sqlite3"
        cls.files = {}
        cls.ids = {}
        with ReferenceCatalog(cls.db_template) as catalog:
            genres = GenrePatternStore(catalog)
            for genre, suffix in (("fantasy", "txt"), ("romance", "md")):
                ids = []
                for index in range(3):
                    source = cls.base / f"synthetic_{genre}_{index}.{suffix}"
                    source.write_text(_synthetic_text(genre, index), encoding="utf-8")
                    imported = import_reference_profile(source, catalog=catalog)
                    ids.append(imported.reference.reference_id)
                    genres.assign(ids[-1], genre)
                    cls.files[(genre, index)] = source
                cls.ids[genre] = tuple(ids)
                genres.build(genre)
                genres.build(genre, weighting="characters")
            ElementSelectionStore(catalog).save("fantasy", ["dialogue_ratio", "emotion_balance"])

    @classmethod
    def tearDownClass(cls):
        cls.root.cleanup()

    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.folder = Path(temp.name)
        self.db = self.folder / "test.sqlite3"
        shutil.copyfile(self.db_template, self.db)
        self.catalog = ReferenceCatalog(self.db)
        self.addCleanup(self.catalog.close)
        self.genres = GenrePatternStore(self.catalog)
        self.profiles = ReferenceProfileStore(self.catalog)
        self.selections = ElementSelectionStore(self.catalog)
        self.strengths = PatternStrengthStore(self.catalog)
        self.usage = PatternUsageStore(self.catalog)

    def _cli(self, main, *arguments):
        out, err = StringIO(), StringIO()
        with redirect_stdout(out), redirect_stderr(err):
            code = main(list(arguments))
        return code, out.getvalue(), err.getvalue()

    def test_imported_cohorts_and_real_profiles(self):
        self.assertEqual(self.catalog.count(), 6)
        self.assertEqual(self.profiles.count(), 6)
        for genre in ("fantasy", "romance"):
            for reference_id in self.ids[genre]:
                record = self.catalog.get(reference_id)
                profile = self.profiles.get(reference_id)
                validate_profile(profile)
                self.assertEqual(record.status, "analyzed")
                self.assertEqual(profile.episode_count, 2)
                self.assertGreater(profile.total_characters, 100)

    def test_source_bytes_never_mutated_by_analysis(self):
        source = self.files[("fantasy", 0)]
        original = source.read_bytes()
        result = import_reference_profile(source, catalog=self.catalog)
        self.assertEqual(source.read_bytes(), original)
        self.assertEqual(result.reference.sha256, hashlib.sha256(original).hexdigest())
        self.assertEqual(result.reference.reference_id, self.ids["fantasy"][0])

    def test_reimport_is_idempotent_and_keeps_genres(self):
        before = self.catalog.count()
        imported = import_reference_profile(self.files[("romance", 1)], catalog=self.catalog)
        self.assertEqual(self.catalog.count(), before)
        self.assertEqual(imported.reference.reference_id, self.ids["romance"][1])
        self.assertEqual(len(self.genres.list_assignments("romance")), 3)

    def test_stored_profile_has_no_source_text_or_author_words(self):
        rows = self.catalog._db.execute("SELECT profile_json FROM reference_profiles").fetchall()
        for row in rows:
            for phrase in (CANARY, "그는 어둠 속에서", "그녀는 꽃을 보고", "제1화"):
                self.assertNotIn(phrase, row["profile_json"])
        self.assertNotIn(CANARY.encode(), self.db.read_bytes())

    def test_rendered_profile_and_pattern_do_not_echo_source(self):
        from ai_novel_factory.references.genre_patterns import render_genre_pattern
        profile = self.profiles.get(self.ids["fantasy"][0])
        output = render_reference_profile(profile) + render_genre_pattern(self.genres.get("fantasy"))
        for forbidden in (CANARY, "지도를 펼쳤다", "마침내 비밀을 발견했다"):
            self.assertNotIn(forbidden, output)

    def test_genre_recipes_for_equal_and_character_weighting(self):
        for genre in ("fantasy", "romance"):
            for weighting in ("equal", "characters"):
                pattern = self.genres.get(genre, weighting=weighting)
                self.assertEqual(pattern.reference_count, 3)
                self.assertEqual(len(pattern.features), 6)
                self.assertEqual(pattern.weighting, weighting)

    def test_selected_features_filter_related_search(self):
        matches = RelatedPatternSearch(self.catalog).search("fantasy").matches
        self.assertEqual(len(matches), 1)
        self.assertEqual(matches[0].genre, "romance")
        self.assertEqual([d.key for d in matches[0].differences],
                         ["dialogue_ratio", "emotion_balance"])

    def test_search_does_not_create_usage_record(self):
        RelatedPatternSearch(self.catalog).search("fantasy")
        self.assertEqual(self.usage.count(), 0)

    def test_usage_record_is_opt_in_and_persists_on_reopen(self):
        result = self.usage.record("fantasy", pattern_genre="romance")
        self.assertEqual(self.usage.count(), 1)
        self.assertTrue(result.current)
        with ReferenceCatalog(self.db) as reopened:
            self.assertEqual(PatternUsageStore(reopened).get(result.usage_id), result)

    def test_usage_history_becomes_stale_after_element_change(self):
        item = self.usage.record("fantasy", pattern_genre="romance")
        self.selections.save("fantasy", ["lexical_diversity"])
        self.assertFalse(self.usage.get(item.usage_id).current)
        self.assertEqual(self.usage.count(), 1)

    def test_strength_bias_blocks_selection_search_and_usage(self):
        for rid in self.ids["fantasy"][1:]:
            self.strengths.set(rid, 0)
        with self.assertRaises(ElementSelectionStorageError):
            self.selections.get("fantasy")
        with self.assertRaises(RelatedPatternError):
            RelatedPatternSearch(self.catalog).search("fantasy")
        with self.assertRaises(PatternUsageError):
            self.usage.record("fantasy")
        self.assertEqual(self.usage.count(), 0)

    def test_strength_restoration_recovers_selection(self):
        excluded = self.ids["fantasy"][1:]
        for rid in excluded:
            self.strengths.set(rid, 0)
        with self.assertRaises(ElementSelectionStorageError):
            self.selections.get("fantasy")
        for rid in excluded:
            self.strengths.reset(rid)
        self.assertEqual(len(self.selections.get("fantasy").selected_keys), 2)

    def test_changed_cohort_requires_rebuild(self):
        rid = self.ids["fantasy"][0]
        self.genres.unassign(rid)
        with self.assertRaises(GenreStorageError):
            self.genres.get("fantasy")
        with self.assertRaises(RelatedPatternError):
            RelatedPatternSearch(self.catalog).search("fantasy")
        self.genres.build("fantasy")
        self.assertEqual(self.genres.get("fantasy").reference_count, 2)

    def test_profile_delete_invalidates_derived_recipe(self):
        self.profiles.delete(self.ids["fantasy"][0])
        with self.assertRaises(GenreStorageError):
            self.genres.get("fantasy")
        self.assertEqual(self.profiles.count(), 5)

    def test_failed_bad_utf8_does_not_register(self):
        source = self.folder / "bad.txt"
        source.write_bytes(b"\x00\x01\xff")
        with self.assertRaises(ReferenceParseError):
            import_reference_profile(source, catalog=self.catalog)
        self.assertEqual(self.catalog.count(), 6)
        self.assertEqual(self.profiles.count(), 6)

    def test_unsupported_epub_fails_closed_without_registration(self):
        source = self.folder / "bad.epub"
        source.write_bytes(b"not a supported parser")
        with self.assertRaises(ReferenceParseError):
            import_reference_profile(source, catalog=self.catalog)
        self.assertEqual(self.catalog.count(), 6)

    def test_size_limit_fails_without_registration(self):
        source = self.folder / "big.txt"
        source.write_text(_synthetic_text("fantasy", 0), encoding="utf8")
        with self.assertRaises(ReferenceParseError):
            import_reference_profile(source, catalog=self.catalog, max_bytes=5)
        self.assertEqual(self.catalog.count(), 6)

    def test_symlink_is_not_imported(self):
        source = self.folder / "private.txt"
        source.write_text(_synthetic_text("fantasy", 0), encoding="utf8")
        link = self.folder / "link.txt"
        try:
            link.symlink_to(source)
        except (OSError, NotImplementedError):
            self.skipTest("symlinks unavailable")
        with self.assertRaises(ReferenceParseError):
            import_reference_profile(link, catalog=self.catalog)
        self.assertEqual(self.catalog.count(), 6)

    def test_network_denied_across_import_and_pattern_search(self):
        source = self.folder / "extra.txt"
        source.write_text(_synthetic_text("fantasy", 0) + " 추가", encoding="utf8")
        with patch.object(socket.socket, "connect", side_effect=AssertionError("network used")), \
             patch.object(socket, "create_connection", side_effect=AssertionError("network used")):
            imported = import_reference_profile(source, catalog=self.catalog)
            self.assertEqual(imported.reference.status, "analyzed")
            self.assertEqual(len(RelatedPatternSearch(self.catalog).search("fantasy").matches), 1)

    def test_cli_uses_same_db_and_reports_numeric_profile(self):
        code, out, err = self._cli(profile_cli, "--database", str(self.db), "genre-show", "fantasy")
        self.assertEqual((code, err), (0, ""))
        self.assertIn("판타지", out)
        self.assertNotIn(CANARY, out)

    def test_cli_can_record_usage_explicitly(self):
        code, out, err = self._cli(profile_cli, "--database", str(self.db), "usage-add", "fantasy")
        self.assertEqual((code, err), (0, ""))
        self.assertEqual(self.usage.count(), 1)
        self.assertNotIn(CANARY, out)

    def test_copy_blocked_from_real_reference_file(self):
        source = self.files[("fantasy", 0)].read_text(encoding="utf8")
        result = analyze_similarity_baseline(source, [source])
        self.assertEqual(result.status, "blocked")
        self.assertFalse(result.allowed_for_automatic_use)
        self.assertNotIn(CANARY, repr(result))

    def test_partial_copy_requires_review_not_automatic_use(self):
        mutated = NEAR_SOURCE.replace("발견하고", "발견하여")
        self.assertEqual(analyze_similarity_baseline(mutated, [NEAR_SOURCE]).status, "review")
        with self.assertRaises(SimilarityCheckError):
            require_similarity_safe_draft(mutated, [NEAR_SOURCE])

    def test_unrelated_text_passes_only_selected_source_baseline(self):
        result = analyze_similarity_baseline(NEAR_SOURCE, [UNRELATED])
        self.assertEqual(result.status, "clear")
        self.assertTrue(result.allowed_for_automatic_use)

    def test_cli_near_copy_uses_real_files_without_excerpts(self):
        ref = self.folder / "reference.txt"
        draft = self.folder / "draft.txt"
        ref.write_text(NEAR_SOURCE + "\n\n" + CANARY, encoding="utf8")
        draft.write_text(NEAR_SOURCE.replace("발견하고", "발견하여"), encoding="utf8")
        code, out, err = self._cli(similarity_cli, "--draft", str(draft), "--reference", str(ref))
        self.assertEqual((code, err), (1, ""))
        self.assertNotIn(CANARY, out)
        self.assertNotIn(NEAR_SOURCE, out)
        self.assertNotIn(str(ref), out)

    def test_cli_missing_reference_fails_without_echoing_path(self):
        draft = self.folder / "draft.txt"
        draft.write_text(NEAR_SOURCE, encoding="utf8")
        missing = self.folder / "PRIVATE_MISSING_PATH_1199.txt"
        code, out, err = self._cli(similarity_cli, "--draft", str(draft), "--reference", str(missing))
        self.assertEqual(code, 2)
        self.assertEqual(out, "")
        self.assertNotIn("PRIVATE_MISSING_PATH_1199", err)

    @unittest.skipUnless(Document is not None, "python-docx optional dependency missing")
    def test_docx_full_import_analysis(self):
        document = Document()
        document.add_heading("제1화", 1)
        document.add_paragraph(_synthetic_text("fantasy", 0).replace("제1화", "").replace("제2화", ""))
        document.add_heading("제2화", 1)
        document.add_paragraph("마지막 문을 열자 새로운 풍경이 펼쳐졌다.")
        source = self.folder / "synthetic.docx"
        document.save(source)
        raw = source.read_bytes()
        result = import_reference_profile(source, catalog=self.catalog)
        self.assertEqual(result.profile.source_format, "docx")
        self.assertEqual(result.profile.episode_count, 2)
        self.assertEqual(source.read_bytes(), raw)
        self.assertNotIn(CANARY, repr(result.profile))

    @unittest.skipUnless(PdfWriter is not None, "pypdf optional dependency missing")
    def test_pdf_full_import_analysis(self):
        writer = PdfWriter()
        page = writer.add_blank_page(width=612, height=792)
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                                 NameObject("/Subtype"): NameObject("/Type1"),
                                 NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
        stream = DecodedStreamObject()
        stream.set_data(b"BT /F1 12 Tf 30 700 Td (Original synthetic PDF content for offline test.) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(stream)
        buffer = BytesIO()
        writer.write(buffer)
        source = self.folder / "synthetic.pdf"
        source.write_bytes(buffer.getvalue())
        result = import_reference_profile(source, catalog=self.catalog)
        self.assertEqual(result.profile.source_format, "pdf")
        self.assertEqual(result.profile.episode_count, 1)
        self.assertGreater(result.profile.total_characters, 0)


if __name__ == "__main__":
    unittest.main()
