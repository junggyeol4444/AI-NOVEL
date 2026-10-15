"""Stage 30 profile unit, SQLite persistence, security and offline CLI tests."""
from __future__ import annotations

from dataclasses import replace
import io
import json
import os
from pathlib import Path
import sqlite3
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

from ai_novel_factory.database import ReferenceCatalog
from ai_novel_factory.database.reference_profiles import ProfileStorageError, ReferenceProfileStore
from ai_novel_factory.references import (
    ReferenceParseError, build_reference_profile, parse_text,
    render_reference_profile,
)
from ai_novel_factory.references.profile_pipeline import import_reference_profile
from ai_novel_factory.references.reference_profile import EpisodeProfile, validate_profile


SAMPLE = (
    "소설 머리말\n\n"
    "제1화 첫 만남\n\n"
    "그는 희망을 느꼈다. 그러나 위험했다.\n\n"
    "'누가 왔지?'라고 물었다.\n\n"
    "제2화 비밀의 끝\n\n"
    "비밀은 결국 드러났다. 사실은 모르는 사람이었다!"
)
SECRET = "DO_NOT_STORE_THIS_ORIGINAL_NOVEL_TEXT_384714"


class ProfileMetricTests(unittest.TestCase):
    def setUp(self):
        self.doc = parse_text(SAMPLE)
        self.profile = build_reference_profile(self.doc)

    def test_profile_contains_fixed_aggregates(self):
        self.assertEqual(self.profile.schema_version, 1)
        self.assertEqual(self.profile.episode_count, 2)
        self.assertGreater(self.profile.total_characters, 0)
        self.assertGreater(self.profile.event_cue_candidates, 0)
        self.assertEqual(len(self.profile.episodes), 2)

    def test_episode_numbers_and_fixed_hook(self):
        self.assertEqual([e.number for e in self.profile.episodes], [1, 2])
        self.assertIn(self.profile.episodes[1].cliffhanger_type, {
            'none', 'question', 'ellipsis', 'exclamation',
            'danger', 'decision', 'interruption', 'revelation',
        })

    def test_dialogue_and_narration_partition(self):
        self.assertAlmostEqual(self.profile.dialogue_ratio + self.profile.narration_ratio, 1, places=4)

    def test_episode_cues_add_up(self):
        self.assertEqual(sum(e.event_cue_candidates for e in self.profile.episodes),
                         self.profile.event_cue_candidates)

    def test_report_is_readable_and_only_fixed_statistics(self):
        report = render_reference_profile(self.profile)
        self.assertIn('참고작 분석 프로필', report)
        self.assertIn('1화', report)
        self.assertNotIn('누가 왔지?', report)
        self.assertNotIn('모르는 사람이었다', report)

    def test_report_limit_and_validation(self):
        out = render_reference_profile(self.profile, episode_limit=1)
        self.assertIn('나머지 1개', out)
        self.assertNotIn('2화 |', out)
        for value in [-1, 10_001, True, 1.5]:
            with self.subTest(value=value), self.assertRaises(ReferenceParseError):
                render_reference_profile(self.profile, episode_limit=value)

    def test_empty_body_headings_can_profile(self):
        doc = parse_text('# 그냥 제목만', format='markdown')
        profile = build_reference_profile(doc)
        self.assertEqual(profile.total_characters, 0)
        self.assertEqual(profile.dialogue_ratio, 0)
        self.assertEqual(profile.narration_ratio, 0)

    def test_text_names_do_not_appear_in_repr(self):
        profile = build_reference_profile(parse_text('제1화\n\n' + SECRET))
        self.assertNotIn(SECRET, repr(profile))
        self.assertNotIn(SECRET, render_reference_profile(profile))

    def test_wrong_profile_datatype(self):
        with self.assertRaises(ReferenceParseError):
            validate_profile({'source_text': SECRET})

    def test_reject_unexpected_format(self):
        with self.assertRaises(ReferenceParseError):
            validate_profile(replace(self.profile, source_format=SECRET))

    def test_reject_injected_type_label(self):
        bad = replace(self.profile.episodes[0], cliffhanger_type=SECRET)
        with self.assertRaises(ReferenceParseError):
            validate_profile(replace(self.profile, episodes=(bad, *self.profile.episodes[1:])))

    def test_reject_nan_inf_bool_ratios(self):
        for value in [float('nan'), float('inf'), float('-inf'), -1, 1.1, True]:
            with self.subTest(value=value), self.assertRaises(ReferenceParseError):
                validate_profile(replace(self.profile, dialogue_ratio=value))

    def test_reject_invalid_count(self):
        for value in [-1, 1.4, True, 10**12]:
            with self.subTest(value=value), self.assertRaises(ReferenceParseError):
                validate_profile(replace(self.profile, total_paragraphs=value))

    def test_reject_wrong_episode_count(self):
        with self.assertRaises(ReferenceParseError):
            validate_profile(replace(self.profile, episode_count=10))

    def test_reject_episode_order(self):
        with self.assertRaises(ReferenceParseError):
            validate_profile(replace(self.profile, episodes=tuple(reversed(self.profile.episodes))))

    def test_reject_episode_aggregation_mismatch(self):
        with self.assertRaises(ReferenceParseError):
            validate_profile(replace(self.profile, event_cue_candidates=0))

    def test_reject_invalid_input_and_stale_split(self):
        with self.assertRaises(ReferenceParseError):
            build_reference_profile('not parsed')
        from ai_novel_factory.references import split_episodes
        with self.assertRaises(ReferenceParseError):
            build_reference_profile(self.doc, split=split_episodes(parse_text('Different text.')))


class ProfileStorageTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.database = self.root / 'local.sqlite3'
        self.catalog = ReferenceCatalog(self.database)
        self.addCleanup(self.catalog.close)
        self.store = ReferenceProfileStore(self.catalog)
        self.f = self.root / 'novel.txt'
        self.f.write_text(SAMPLE, encoding='utf-8')

    def test_import_profile_persists_and_listed(self):
        result = import_reference_profile(self.f, catalog=self.catalog)
        self.assertEqual(result.reference.status, 'analyzed')
        self.assertEqual(self.store.count(), 1)
        self.assertEqual(self.store.get(result.reference.reference_id), result.profile)
        self.assertEqual(self.store.list()[0].file_name, 'novel.txt')

    def test_profile_survives_reopen_without_source_file(self):
        item = import_reference_profile(self.f, catalog=self.catalog)
        self.f.unlink()
        with ReferenceCatalog(self.database) as reopened:
            loaded = ReferenceProfileStore(reopened).get(item.reference.reference_id)
            self.assertEqual(loaded, item.profile)

    def test_idempotent_duplicate_import(self):
        a = import_reference_profile(self.f, catalog=self.catalog)
        b = import_reference_profile(self.f, catalog=self.catalog)
        self.assertEqual(a.reference.reference_id, b.reference.reference_id)
        self.assertEqual(self.store.count(), 1)
        self.assertEqual(self.catalog.count(), 1)

    def test_corrupt_input_leaves_db_clean(self):
        path = self.root / 'broken.txt'
        path.write_bytes(b'\x00broken')
        with self.assertRaises(ReferenceParseError):
            import_reference_profile(path, catalog=self.catalog)
        self.assertEqual(self.catalog.count(), 0)
        self.assertEqual(self.store.count(), 0)

    def test_invalid_epub_fails_closed(self):
        # EPUB is supported now; malformed content must still be rejected
        # before a catalog row or profile is persisted.
        source = self.root / 'test.epub'
        source.write_bytes(b'synthetic invalid EPUB bytes')
        with self.assertRaises(ReferenceParseError):
            import_reference_profile(source, catalog=self.catalog)
        self.assertEqual(self.catalog.count(), 0)
        self.assertEqual(self.store.count(), 0)

    def test_database_has_no_raw_source_and_paths(self):
        self.f.write_text('제1화\n\n' + SECRET, encoding='utf-8')
        result = import_reference_profile(self.f, catalog=self.catalog)
        self.assertNotIn(SECRET, repr(result))
        dump = '\n'.join(self.catalog._db.iterdump())
        self.assertNotIn(SECRET, dump)
        self.assertNotIn(str(self.root), dump)
        columns = {r[1] for r in self.catalog._db.execute('PRAGMA table_info(reference_profiles)')}
        self.assertEqual(columns, {'reference_id','schema_version','profile_json','updated_at'})

    def test_manual_profile_save_requires_registered_file(self):
        prof = build_reference_profile(parse_text(SAMPLE))
        with self.assertRaises(ProfileStorageError):
            self.store.save('nonexistent-id', prof)

    def test_invalid_profile_is_not_persisted(self):
        result = import_reference_profile(self.f, catalog=self.catalog)
        before = self.store.get(result.reference.reference_id)
        with self.assertRaises(ReferenceParseError):
            self.store.save(result.reference.reference_id, replace(before, source_format=SECRET))
        self.assertEqual(self.store.get(result.reference.reference_id), before)

    def test_profiles_cascade_when_metadata_is_explicitly_removed(self):
        result = import_reference_profile(self.f, catalog=self.catalog)
        self.assertEqual(self.catalog._db.execute("PRAGMA foreign_keys").fetchone()[0], 1)
        with self.catalog._db:
            self.catalog._db.execute(
                'DELETE FROM reference_files WHERE reference_id=?',
                (result.reference.reference_id,),
            )
        self.assertEqual(self.store.count(), 0)

    def test_profile_format_must_match_registered_file(self):
        result = import_reference_profile(self.f, catalog=self.catalog)
        with self.assertRaises(ProfileStorageError):
            self.store.save(result.reference.reference_id,
                            replace(result.profile, source_format='pdf'))

    def test_no_profile_for_basic_stage20_import(self):
        from ai_novel_factory.references import import_reference_file
        record = import_reference_file(self.f, catalog=self.catalog).reference
        with self.assertRaises(ProfileStorageError):
            self.store.get(record.reference_id)
        self.assertEqual(self.store.count(), 0)

    def test_list_pagination_and_delete_only_derived_metrics(self):
        r = import_reference_profile(self.f, catalog=self.catalog)
        self.assertEqual(len(self.store.list(limit=1, offset=0)), 1)
        self.assertEqual(self.store.list(limit=1, offset=1), ())
        self.assertTrue(self.store.delete(r.reference.reference_id))
        self.assertFalse(self.store.delete(r.reference.reference_id))
        self.assertEqual(self.store.count(), 0)
        self.assertEqual(self.catalog.count(), 1)

    def test_list_invalid_pagination(self):
        for kwargs in ({'limit': True}, {'limit':0}, {'limit':1001},
                       {'offset': -1}, {'offset':1.5}):
            with self.subTest(kwargs=kwargs), self.assertRaises(ProfileStorageError):
                self.store.list(**kwargs)

    def test_corrupted_json_rejected_safely(self):
        ref = import_reference_profile(self.f, catalog=self.catalog).reference
        self.catalog._db.execute('UPDATE reference_profiles SET profile_json=? WHERE reference_id=?',
                                 ('{"source_text":"' + SECRET + '"}', ref.reference_id))
        self.catalog._db.commit()
        with self.assertRaises(ProfileStorageError) as cm:
            self.store.get(ref.reference_id)
        self.assertNotIn(SECRET, str(cm.exception))

    def test_old_schema_is_rejected(self):
        ref = import_reference_profile(self.f, catalog=self.catalog).reference
        with self.assertRaises(sqlite3.IntegrityError):
            self.catalog._db.execute('UPDATE reference_profiles SET schema_version=9 WHERE reference_id=?',
                                     (ref.reference_id,))
        self.catalog._db.rollback()

    def test_id_lookup_is_parameterized(self):
        ref = import_reference_profile(self.f, catalog=self.catalog).reference
        self.assertEqual(self.store.count(), 1)
        with self.assertRaises(ProfileStorageError):
            self.store.get(ref.reference_id + "' OR '1'='1")
        self.assertEqual(self.store.count(), 1)

    def test_no_network_used_in_import(self):
        with patch('socket.socket', side_effect=AssertionError('network called')):
            result = import_reference_profile(self.f, catalog=self.catalog)
        self.assertEqual(self.store.get(result.reference.reference_id), result.profile)

    def test_permissionless_no_implicit_directory_creation(self):
        import_reference_profile(self.f, catalog=self.catalog)
        self.assertFalse((self.root / 'downloaded-novels').exists())

    def test_invalid_import_limit(self):
        for value in [True, 0, 3.5]:
            with self.subTest(value=value), self.assertRaises(ReferenceParseError):
                import_reference_profile(self.f, catalog=self.catalog, max_bytes=value)

    def test_no_absolute_path_in_safe_errors(self):
        bad = self.root / 'private_credentials_password.txt'
        with self.assertRaises(ReferenceParseError) as cm:
            import_reference_profile(bad, catalog=self.catalog)
        self.assertNotIn('private_credentials_password', str(cm.exception))


class ProfileCLITests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.db = self.root / 'local.sqlite3'
        self.book = self.root / 'fiction.txt'
        self.book.write_text(SAMPLE, encoding='utf-8')

    def cli(self, *args):
        from ai_novel_factory.references.profile_cli import main
        with patch('sys.stdout', new_callable=io.StringIO) as stdout, patch('sys.stderr', new_callable=io.StringIO) as stderr:
            rc = main(['--database', str(self.db), *args])
        return rc, stdout.getvalue(), stderr.getvalue()

    def test_cli_import_and_list_and_show(self):
        code, report, error = self.cli('import', str(self.book))
        self.assertEqual((code,error), (0,''))
        self.assertIn('등록 ID:', report)
        code, listing, error = self.cli('list')
        self.assertEqual((code,error), (0,''))
        self.assertIn('fiction.txt', listing)
        ref_id = listing.split()[0]
        code, detail, error = self.cli('show', ref_id)
        self.assertEqual((code,error), (0,''))
        self.assertIn('참고작 분석 프로필', detail)
        self.assertIn('fiction.txt', detail)

    def test_cli_error_does_not_expose_file_path(self):
        code, out, err = self.cli('import', str(self.root/'private_secret_book.txt'))
        self.assertEqual(code, 1)
        self.assertNotIn('private_secret_book', err)

    def test_cli_explicit_db_parent_required(self):
        db = self.root / 'missing' / 'db.sqlite3'
        from ai_novel_factory.references.profile_cli import main
        with patch('sys.stdout', new_callable=io.StringIO), patch('sys.stderr', new_callable=io.StringIO) as err:
            code = main(['--database', str(db), 'list'])
        self.assertEqual(code, 1)
        self.assertNotIn(str(self.root), err.getvalue())

    def test_cli_subprocess_real_entrypoint(self):
        env = dict(os.environ)
        env['PYTHONPATH'] = str(Path(__file__).resolve().parents[1] / 'src')
        command = [sys.executable, '-m', 'ai_novel_factory.references.profile_cli',
                   '--database', str(self.db), 'import', str(self.book)]
        proc = subprocess.run(command, capture_output=True, text=True, env=env, timeout=30)
        self.assertEqual(proc.returncode, 0, proc.stderr)
        self.assertIn('등록 ID:', proc.stdout)
        cmd2 = [sys.executable, '-m', 'ai_novel_factory.references.profile_cli',
                '--database', str(self.db), 'list']
        output = subprocess.run(cmd2, capture_output=True, text=True, env=env, timeout=20)
        self.assertEqual(output.returncode, 0, output.stderr)
        self.assertIn('fiction.txt', output.stdout)


if __name__ == '__main__':
    unittest.main()

class ProfileFormatIntegrationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.catalog = ReferenceCatalog(self.root / 'formats.db')
        self.addCleanup(self.catalog.close)

    def test_markdown_profile_uses_md_catalog_name(self):
        file = self.root / 'test.markdown'
        file.write_text('# 제1화\n\n예언이 있었다.\n\n# 제2화\n\n진실이 드러났다.', encoding='utf-8')
        result = import_reference_profile(file, catalog=self.catalog)
        self.assertEqual(result.reference.source_format, 'md')
        self.assertEqual(result.profile.source_format, 'markdown')
        self.assertEqual(ReferenceProfileStore(self.catalog).get(result.reference.reference_id), result.profile)

    def test_docx_integration(self):
        try:
            from docx import Document
        except ImportError:
            self.skipTest('python-docx missing')
        from io import BytesIO
        doc = Document()
        doc.add_heading('제1화', level=1)
        doc.add_paragraph('희망이 보였다.')
        doc.add_heading('제2화', level=1)
        doc.add_paragraph('그러나 위험했다.')
        file = self.root / 'test.docx'
        doc.save(file)
        result = import_reference_profile(file, catalog=self.catalog)
        self.assertEqual(result.reference.source_format, 'docx')
        self.assertEqual(result.profile.episode_count, 2)

    def test_pdf_integration(self):
        try:
            from tests.test_import_pipeline import sample_pdf
        except ImportError:
            try:
                from test_import_pipeline import sample_pdf
            except ImportError:
                self.skipTest('pypdf sample unavailable')
        file = self.root / 'test.pdf'
        file.write_bytes(sample_pdf())
        result = import_reference_profile(file, catalog=self.catalog)
        self.assertEqual(result.reference.source_format, 'pdf')
        self.assertEqual(result.profile.episode_count, 1)

    def test_profiles_remain_offline_when_internet_unavailable(self):
        file = self.root / 'offline.md'
        file.write_text('# 제1화\n\n아무도 접속하지 않았다.', encoding='utf-8')
        with patch('urllib.request.urlopen', side_effect=AssertionError('network')):
            with patch('socket.socket', side_effect=AssertionError('network')):
                result = import_reference_profile(file, catalog=self.catalog)
        self.assertEqual(ReferenceProfileStore(self.catalog).get(result.reference.reference_id), result.profile)
